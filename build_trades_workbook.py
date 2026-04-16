"""Parse BL Final Table transaction transcripts and emit an Excel workbook.

Usage:
    python build_trades_workbook.py trades_raw.txt BL_Trades.xlsx
"""

import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TRADE_RE = re.compile(
    r"^(?P<date>[A-Za-z]{3,} \d+)(?: - [^:]+)?:\s+"
    r"Team\s+\d+\s+\((?P<a>[^)]+)\)\s+reported trading"
    r"(?:\s+(?P<ags>.*?))?\s+to\s+team\s+\d+\s+\((?P<b>[^)]+)\)"
    r"(?:\s+for\s+(?P<bgs>.*?))?\s*\.\s*$"
)


def split_block_years(lines):
    """Return list of (year_idx, line) tagging each line with a 1-based block index."""
    out = []
    year = 0
    for line in lines:
        if line.startswith("BL Final Table Transactions"):
            year += 1
        out.append((max(year, 1), line))
    return out


def first_initial_last(full_name: str) -> str:
    """FirstName + first initial of last name. Handles middle initials/names."""
    parts = full_name.replace(".", "").split()
    if len(parts) == 1:
        return parts[0]
    first = parts[0]
    last = parts[-1]
    return f"{first}{last[0]}"


def sheet_safe(name: str) -> str:
    return re.sub(r"[\\/*?:\[\]]", "", name)[:31]


def normalize_items(raw: str | None) -> tuple[str, ...]:
    if not raw:
        return ()
    # Split on "; " and normalize whitespace
    parts = [p.strip() for p in raw.split(";")]
    return tuple(sorted(p for p in parts if p))


def parse(path: Path):
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = text.splitlines()
    tagged = split_block_years(lines)
    trades = []  # list of dicts
    for year_idx, line in tagged:
        m = TRADE_RE.match(line.strip())
        if not m:
            continue
        a = m.group("a").strip()
        b = m.group("b").strip()
        ags = (m.group("ags") or "").strip()
        bgs = (m.group("bgs") or "").strip()
        date = m.group("date").strip()
        trades.append(
            {
                "year": year_idx,
                "date": date,
                "from": a,
                "to": b,
                "from_gave": ags,
                "from_got": bgs,
                "raw": line.strip(),
            }
        )
    return trades


def canonical_key(t):
    """Key that matches both-sides reports of the same trade."""
    a, b = t["from"], t["to"]
    ags = normalize_items(t["from_gave"])
    bgs = normalize_items(t["from_got"])
    pair = tuple(sorted([a, b]))
    # Order items to match owner-sort order
    if a <= b:
        items = (ags, bgs)
    else:
        items = (bgs, ags)
    return (t["year"], t["date"], pair, items)


def trade_summary(t: dict) -> str:
    """Readable one-line description of a trade for display in a cell."""
    ags = t["from_gave"] or "(nothing listed)"
    bgs = t["from_got"] or "(nothing listed)"
    return f"{t['date']}: {t['from']} sent {ags} -> {t['to']} for {bgs}"


def build_workbook(trades, out_path: Path):
    wb = Workbook()

    # --- All Trades sheet ---
    ws_all = wb.active
    ws_all.title = "All Trades"
    headers = ["Year Block", "Date", "From Owner", "To Owner", "From Gave", "From Got", "Full Line"]
    ws_all.append(headers)
    for cell in ws_all[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for t in trades:
        ws_all.append([
            t["year"],
            t["date"],
            t["from"],
            t["to"],
            t["from_gave"],
            t["from_got"],
            t["raw"],
        ])
    widths = [10, 10, 22, 22, 40, 40, 100]
    for i, w in enumerate(widths, 1):
        ws_all.column_dimensions[get_column_letter(i)].width = w
    ws_all.freeze_panes = "A2"

    # --- De-dup into unique trades ---
    unique = {}
    for t in trades:
        k = canonical_key(t)
        # Keep whichever direction came first; doesn't matter which
        if k not in unique:
            unique[k] = t

    # owner -> list of (partner, trade) for unique trades only
    by_owner_partner = defaultdict(lambda: defaultdict(list))
    all_owners = set()
    for t in unique.values():
        a, b = t["from"], t["to"]
        all_owners.add(a)
        all_owners.add(b)
        by_owner_partner[a][b].append(t)
        by_owner_partner[b][a].append(t)

    owners_sorted = sorted(all_owners)

    # Sheet-name collision handling: ensure unique short names
    used_titles = {"All Trades"}

    def unique_title(base):
        t = sheet_safe(base)
        cand = t
        i = 2
        while cand in used_titles:
            cand = sheet_safe(f"{t}{i}")
            i += 1
        used_titles.add(cand)
        return cand

    for owner in owners_sorted:
        title = unique_title(first_initial_last(owner))
        ws = wb.create_sheet(title)
        ws["A1"] = f"{owner} - trades with each other owner"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:D1")

        ws.append([])  # blank row
        header_row = ["Other Owner", "# Trades"]
        # Figure max trades so we know how many trade columns
        partners = sorted(p for p in by_owner_partner[owner].keys())
        max_trades = max((len(by_owner_partner[owner][p]) for p in partners), default=0)
        for i in range(1, max_trades + 1):
            header_row.append(f"Trade {i}")
        ws.append(header_row)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")

        total = 0
        for p in partners:
            ts = sorted(by_owner_partner[owner][p], key=lambda x: (x["year"], x["date"]))
            row = [p, len(ts)]
            total += len(ts)
            for t in ts:
                row.append(trade_summary(t))
            ws.append(row)

        # Total row
        total_row = ["TOTAL", total]
        ws.append(total_row)
        for cell in ws[ws.max_row][:2]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FFF2CC")

        # Column widths
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 10
        for i in range(3, 3 + max_trades):
            ws.column_dimensions[get_column_letter(i)].width = 80
        ws.freeze_panes = "C4"

    wb.save(out_path)


def main():
    if len(sys.argv) < 3:
        print("Usage: build_trades_workbook.py INPUT.txt OUTPUT.xlsx")
        sys.exit(1)
    inp = Path(sys.argv[1])
    out = Path(sys.argv[2])
    trades = parse(inp)
    print(f"Parsed {len(trades)} trade lines.")
    build_workbook(trades, out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
